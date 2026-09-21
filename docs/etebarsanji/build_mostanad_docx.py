#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build RTL Word document from docs/مستند-اعتبارسنجی.md."""

from __future__ import annotations

import re
import sys
from pathlib import Path

# Reuse RTL styling from sibling script
sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_support_brief import (  # noqa: E402
    BODY_FONT,
    GRAY,
    NAVY,
    add_heading_custom,
    add_p,
    add_table,
    set_run,
    set_paragraph_rtl,
)
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "مستند-اعتبارسنجی.md"
OUT_REPO = ROOT / "مستند-اعتبارسنجی.docx"
OUT_ART = Path("/opt/cursor/artifacts/مستند-اعتبارسنجی.docx")


def strip_md_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    return text.strip()


def add_bullet(doc, text: str, *, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_rtl(p, align="right", space_after=4, space_before=0, line=1.15)
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.4)
    run = p.add_run()
    set_run(run, strip_md_inline(text), font=BODY_FONT, size=11.5)


def parse_table_row(line: str) -> list[str] | None:
    line = line.strip()
    if not line.startswith("|"):
        return None
    parts = [c.strip() for c in line.strip("|").split("|")]
    return parts


def is_separator_row(cells: list[str]) -> bool:
    return all(re.match(r"^:?-+:?$", c.replace(" ", "")) for c in cells)


def build_doc(md_text: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    lines = md_text.splitlines()
    i = 0
    table_buffer: list[list[str]] | None = None

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("|"):
            row = parse_table_row(stripped)
            if row is None:
                i += 1
                continue
            if is_separator_row(row):
                i += 1
                continue
            if table_buffer is None:
                table_buffer = [row]
            else:
                table_buffer.append(row)
            i += 1
            # flush table when next line is not table
            if i < len(lines) and not lines[i].strip().startswith("|"):
                headers = [strip_md_inline(c) for c in table_buffer[0]]
                rows = [[strip_md_inline(c) for c in r] for r in table_buffer[1:]]
                add_table(doc, headers, rows)
                table_buffer = None
            continue

        if table_buffer is not None:
            headers = [strip_md_inline(c) for c in table_buffer[0]]
            rows = [[strip_md_inline(c) for c in r] for r in table_buffer[1:]]
            add_table(doc, headers, rows)
            table_buffer = None

        if stripped.startswith("# "):
            add_heading_custom(doc, strip_md_inline(stripped[2:]), level=0)
            i += 1
            continue
        if stripped.startswith("## "):
            add_heading_custom(doc, strip_md_inline(stripped[3:]), level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            add_heading_custom(doc, strip_md_inline(stripped[4:]), level=2)
            i += 1
            continue

        if stripped.startswith("- "):
            add_bullet(doc, stripped[2:])
            i += 1
            continue

        if stripped.startswith("**") and stripped.endswith("**") and len(stripped) > 4:
            add_p(doc, strip_md_inline(stripped), size=11.5, bold=True, color=NAVY, align="right")
            i += 1
            continue

        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            add_p(doc, strip_md_inline(stripped), size=10.5, color=GRAY, align="right")
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        add_p(doc, strip_md_inline(stripped), size=11.5, align="justify")
        i += 1

    if table_buffer is not None:
        headers = [strip_md_inline(c) for c in table_buffer[0]]
        rows = [[strip_md_inline(c) for c in r] for r in table_buffer[1:]]
        add_table(doc, headers, rows)

    return doc


def main() -> None:
    if not MD_PATH.is_file():
        raise SystemExit(f"Markdown not found: {MD_PATH}")
    md_text = MD_PATH.read_text(encoding="utf-8")
    doc = build_doc(md_text)
    OUT_REPO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_REPO)
    OUT_ART.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_ART)
    print(f"Wrote {OUT_REPO}")
    print(f"Wrote {OUT_ART}")


if __name__ == "__main__":
    main()
