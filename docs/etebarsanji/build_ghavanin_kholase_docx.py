#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""خلاصهٔ قابل‌فهم گزارش قوانین کد — جایگزین نسخهٔ فنی سنگین."""

from __future__ import annotations

import importlib.util
import re
import sys
from pathlib import Path

_DIR = Path(__file__).resolve().parent

# Word RTL helpers
_spec_sb = importlib.util.spec_from_file_location("build_support_brief", _DIR / "build_support_brief.py")
_sb = importlib.util.module_from_spec(_spec_sb)
_spec_sb.loader.exec_module(_sb)

# Parser + section blurbs
_spec_brief = importlib.util.spec_from_file_location("build_code_laws_brief", _DIR / "build_code_laws_brief.py")
_brief = importlib.util.module_from_spec(_spec_brief)
_spec_brief.loader.exec_module(_brief)

# تابع / شرح بدون شناسه
_spec_md = importlib.util.spec_from_file_location("build_mostanad_md", _DIR / "build_mostanad_md.py")
_md = importlib.util.module_from_spec(_spec_md)
_spec_md.loader.exec_module(_md)

add_p = _sb.add_p
add_heading_custom = _sb.add_heading_custom
add_table = _sb.add_table
TEAL = _sb.TEAL
NAVY = _sb.NAVY
GRAY = _sb.GRAY
HEAD_FONT = _sb.HEAD_FONT

parse_file = _brief.parse_file
FILES = _brief.FILES
RULES = _brief.RULES
SECTION_TITLE = _brief.SECTION_TITLE
SECTION_BLURB = _brief.SECTION_BLURB

pick_applier = _md.pick_applier
law_description = _md.law_description
collapse_duplicate_phrases = _md.collapse_duplicate_phrases
replace_workflows = _brief.replace_workflows
MAIN_FUNC = _md.MAIN_FUNC
_file_key = _md._file_key

OUT_PATHS = [
    _DIR / "Gozaresh-Ghavanin-Kod-Etebarsanji.docx",
    _DIR / "گزارش-قوانین-کد-اعتبارسنجی-شهرسازی.docx",
    Path("/opt/cursor/artifacts/Gozaresh-Ghavanin-Kod-Etebarsanji.docx"),
    Path("/opt/cursor/artifacts/گزارش-قوانین-کد-اعتبارسنجی-شهرسازی.docx"),
]

JARGON = [
    (r"\bdead\s*code\b", "کد غیرفعال"),
    (r"\bExit Function\b", "خروج از تابع"),
    (r"\binline\b", ""),
    (r"`[^`]+`", ""),
    (r"\bIf\s+1\s*=\s*2\b", "شرط همیشه خاموش"),
    (r"\bMsgbox\b", "پنجره پیام"),
    (r"\btmpStr\b", "شناسه پرونده"),
    (r"\bNidProc\b", "شناسه پرونده"),
    (r"\bFormName\b", "نام فرم"),
    (r"\bGUID\b", ""),
    (r"\bSP\s+\d+", ""),
    (r"\bSAMPA\s+\d+", ""),
    (r"\bRE\s+\d+", ""),
    (r"\s+—\s+—", " —"),
]


def layman(text: str, max_len: int = 220) -> str:
    if not text:
        return ""
    t = replace_workflows(text)
    t = collapse_duplicate_phrases(t)
    for pat, repl in JARGON:
        t = re.sub(pat, repl, t, flags=re.I)
    t = re.sub(r"\s+", " ", t).strip(" —-،.")
    if len(t) > max_len:
        t = t[: max_len - 1].rsplit(" ", 1)[0] + "…"
    return t


def setup_doc():
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    for section in doc.sections:
        section.page_width = __import__("docx.shared", fromlist=["Cm"]).Cm(21.0)
        section.page_height = __import__("docx.shared", fromlist=["Cm"]).Cm(29.7)
        section.top_margin = section.bottom_margin = __import__("docx.shared", fromlist=["Cm"]).Cm(1.8)
        section.right_margin = __import__("docx.shared", fromlist=["Cm"]).Cm(2.0)
        section.left_margin = __import__("docx.shared", fromlist=["Cm"]).Cm(1.7)
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        section._sectPr.append(bidi)
    return doc


def build() -> int:
    doc = setup_doc()
    add_p(doc, "بسمه تعالی", size=13, bold=True, color=NAVY, align="center", space_after=6)
    add_heading_custom(doc, "گزارش خلاصه قوانین اعتبارسنجی در کد", level=0)
    add_p(
        doc,
        "این سند نسخهٔ خلاصه و قابل‌فهم گزارش فنی است. "
        "به‌جای جزئیات برنامه‌نویسی، می‌گوید هر فرم چه می‌کند و هنگام ذخیره چه کنترل‌هایی اجرا می‌شود. "
        "نام نوع درخواست با عنوان فارسی فرایند آمده است، نه شناسهٔ فنی.",
        size=11.5,
        color=TEAL,
        align="center",
        space_after=14,
    )

    add_heading_custom(doc, "۱. قبل از هر چیز", level=1)
    add_table(
        doc,
        ["واژه", "یعنی چه برای کارشناس"],
        [
            ["ذخیره", "کارشناس دکمه ذخیره را می‌زند؛ برنامه کنترل‌ها را یکی‌یکی اجرا می‌کند."],
            ["توقف (Stop)", "پیام خطا؛ ذخیره انجام نمی‌شود."],
            ["هشدار (Warning)", "پیام هشدار؛ معمولاً ذخیره ادامه می‌یابد."],
            ["تأیید مدیر", "اگر مدیر قبلاً تأیید کرده باشد، بعضی فیلدها قفل می‌شوند."],
            ["فیش درآمد", "اگر فیش تأیید شده باشد، ویرایش بروکف، بازدید، توافق و … سخت‌تر است."],
            ["ثامن", "ناحیهٔ ۸۰؛ بسیاری از قوانین فقط برای ثامن یا فقط غیر ثامن است."],
            ["نوع درخواست", "مثلاً پروانه، مفاصا، گواهی عدم خلاف — با عنوان فارسی در جدول‌ها."],
        ],
    )
    add_p(
        doc,
        "مسیر معمول پرونده: ثبت درخواست → بازدید → بروکف → پرونده → تحلیل تخلف → تأیید مدیر → "
        "صلح / توافق / تخفیف / جریمه / ضابطه → مأمور بازدید → درآمد → موافقت اصولی → شهروندسپاری.",
        first_line=0.4,
        space_after=12,
    )

    add_heading_custom(doc, "۲. فرم‌ها و کنترل‌های ذخیره", level=1)
    total = 0
    summary_rows = []

    for _pnum, fname in FILES:
        key = _file_key(fname)
        data = parse_file(RULES / fname)
        title = SECTION_TITLE.get(fname, data["title"])
        default_fn = MAIN_FUNC[key]
        n = len(data["laws"])
        total += n
        summary_rows.append([title, str(n)])

        add_heading_custom(doc, title, level=1)
        for blurb in SECTION_BLURB.get(fname, []):
            add_p(doc, layman(blurb, max_len=500), size=11.5, first_line=0.35)
        add_p(
            doc,
            f"تابع اصلی در برنامه: {default_fn}()",
            size=10.5,
            color=GRAY,
            align="right",
            space_after=8,
        )

        rows = []
        for law in data["laws"]:
            fn = pick_applier(law, default_fn)
            desc = layman(law_description(law))
            if desc:
                rows.append([f"{fn}()", desc])

        if rows:
            add_table(doc, ["تابع اعمال‌کننده", "خلاصهٔ کنترل"], rows)
        else:
            add_p(doc, "کنترل استخراج‌نشده.", color=GRAY)

    add_heading_custom(doc, "۳. جمع‌بندی", level=1)
    add_table(doc, ["فرم", "تعداد کنترل هنگام ذخیره"], summary_rows + [["جمع", str(total)]])
    add_p(
        doc,
        "اگر پیام روی صفحه با این جدول فرق داشت، متن همان پیام روی صفحه ملاک است. "
        "برای شرط دقیق و استثناها به کد فرمول یا فایل‌های قانون در مخزن مراجعه کنید.",
        size=11,
        color=GRAY,
        first_line=0.4,
    )

    first = OUT_PATHS[0]
    first.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(first))
    blob = first.read_bytes()
    for p in OUT_PATHS[1:]:
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(blob)
    print("Wrote", first, "size", first.stat().st_size, "laws", total)
    return total


if __name__ == "__main__":
    sys.exit(0 if build() else 1)
