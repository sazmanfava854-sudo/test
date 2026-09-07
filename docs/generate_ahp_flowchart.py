#!/usr/bin/env python3
"""Generate AHP workflow flowchart PowerPoint (light theme matching reference image)."""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUTPUT_DIR = Path("/workspace/docs")
PPTX_PATH = OUTPUT_DIR / "AHP_Workflow_Flowchart.pptx"
PNG_PATH = OUTPUT_DIR / "AHP_Workflow_Flowchart.png"

# Palette matched to reference flowchart (light background, blue / teal / navy)
C_BLUE = (37, 99, 210)
C_LIGHT_BLUE = (96, 165, 250)
C_TEAL = (14, 165, 170)
C_NAVY = (30, 58, 138)
C_GREEN = (34, 197, 94)

STEPS = [
    {
        "num": "۱",
        "title": "گردآوری نظرات خبرگان",
        "subtitle": "جمع‌آوری ماتریس‌های مقایسه زوجی",
        "position": "top",
        "start_label": "شروع",
        "circle_style": "outline_blue",
        "box_border": C_BLUE,
    },
    {
        "num": "۲",
        "title": "مقایسه‌های زوجی",
        "subtitle": "با AHP",
        "position": "bottom",
        "circle_style": "filled_light_blue",
        "box_border": C_LIGHT_BLUE,
    },
    {
        "num": "۳",
        "title": "بررسی سازگاری",
        "subtitle": "محاسبه CR برای هر خبره",
        "position": "top",
        "circle_style": "outline_teal",
        "box_border": C_TEAL,
    },
    {
        "num": "۵",
        "title": "میانگین هندسی",
        "subtitle": "تجمیع نظرات خبرگان",
        "position": "top",
        "connector_label": "گردآوری",
        "connector_color": C_NAVY,
        "circle_style": "filled_light_blue",
        "box_border": C_LIGHT_BLUE,
    },
    {
        "num": "۶",
        "title": "بررسی سازگاری ماتریس تجمیع‌شده",
        "subtitle": "محاسبه CR ماتریس نهایی",
        "position": "bottom",
        "box_w": 300,
        "box_h": 118,
        "circle_style": "outline_teal",
        "box_border": C_TEAL,
    },
    {
        "num": "۷",
        "title": "وزن نهایی معیارها",
        "subtitle": "ماتریس تجمیع‌شده مبنای محاسبه",
        "position": "bottom",
        "final": True,
        "circle_style": "outline_navy",
        "box_border": C_NAVY,
    },
]

BG = (255, 255, 255)
LINE_TOP = (210, 218, 230)
LINE_BOTTOM = C_BLUE
BOX_BG = (255, 255, 255)
FINAL_BG = C_NAVY
TEXT_DARK = (30, 41, 59)
TEXT_MUTED = (100, 116, 139)
TEXT_WHITE = (255, 255, 255)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/workspace/docs/fonts/Vazirmatn-Bold.ttf",
        "/workspace/docs/fonts/Vazirmatn-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    if bold:
        candidates = [c for c in candidates if "Bold" in c] + [c for c in candidates if "Bold" not in c]
    for path in candidates:
        p = Path(path)
        if p.exists():
            return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()


def text_size(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont) -> tuple[int, int]:
    bbox = draw.textbbox((0, 0), text, font=font)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    words = text.split()
    if not words:
        return [text]
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        trial = f"{current} {word}"
        if text_size(draw, trial, font)[0] <= max_width:
            current = trial
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def draw_centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    lines: list[tuple[str, ImageFont.FreeTypeFont, tuple[int, int, int]]],
    gap: int = 8,
) -> None:
    x0, y0, x1, y1 = box
    max_w = x1 - x0 - 24
    rendered: list[tuple[str, ImageFont.FreeTypeFont, tuple[int, int, int], int]] = []
    for text, font, color in lines:
        for chunk in wrap_text(draw, text, font, max_w):
            _, h = text_size(draw, chunk, font)
            rendered.append((chunk, font, color, h))
    total_h = sum(h for *_, h in rendered) + gap * (len(rendered) - 1)
    y = y0 + max(8, (y1 - y0 - total_h) // 2)
    for chunk, font, color, h in rendered:
        w, _ = text_size(draw, chunk, font)
        x = x0 + (x1 - x0 - w) // 2
        draw.text((x, y), chunk, font=font, fill=color)
        y += h + gap


def draw_circle_node(
    draw: ImageDraw.ImageDraw,
    x: int,
    y: int,
    r: int,
    num: str,
    style: str,
    num_font: ImageFont.FreeTypeFont,
) -> None:
    bbox = (x - r, y - r, x + r, y + r)
    border = 4

    if style == "filled_light_blue":
        draw.ellipse(bbox, fill=C_LIGHT_BLUE, outline=C_LIGHT_BLUE, width=border)
        num_color = TEXT_WHITE
    elif style == "outline_teal":
        draw.ellipse(bbox, fill=BG, outline=C_TEAL, width=border)
        num_color = C_TEAL
    elif style == "outline_navy":
        draw.ellipse(bbox, fill=BG, outline=C_NAVY, width=border)
        num_color = C_NAVY
    else:  # outline_blue
        draw.ellipse(bbox, fill=BG, outline=C_BLUE, width=border)
        num_color = C_BLUE

    nw, nh = text_size(draw, num, num_font)
    draw.text((x - nw // 2, y - nh // 2 - 2), num, font=num_font, fill=num_color)


def render_flowchart_png(path: Path) -> None:
    width, height = 1920, 720
    img = Image.new("RGB", (width, height), BG)
    draw = ImageDraw.Draw(img)

    title_font = load_font(34, bold=True)
    sub_font = load_font(22)
    num_font = load_font(28, bold=True)
    label_font = load_font(18)
    small_font = load_font(16)

    margin_x = 120
    line_y = height // 2
    n = len(STEPS)
    spacing = (width - 2 * margin_x) / (n - 1)
    xs = [int(width - margin_x - i * spacing) for i in range(n)]

    # Two-tone timeline (grey top + blue bottom)
    draw.line([(margin_x, line_y - 3), (width - margin_x, line_y - 3)], fill=LINE_TOP, width=8)
    draw.line([(margin_x, line_y + 3), (width - margin_x, line_y + 3)], fill=LINE_BOTTOM, width=8)

    circle_r = 34
    default_box_w, default_box_h = 270, 110

    for i, step in enumerate(STEPS):
        x = xs[i]
        is_final = step.get("final", False)
        box_w = step.get("box_w", default_box_w)
        box_h = step.get("box_h", default_box_h)
        box_border = step["box_border"]

        if i > 0 and step.get("connector_label"):
            prev_x = xs[i - 1]
            label = step["connector_label"]
            color = step.get("connector_color", C_NAVY)
            lw, lh = text_size(draw, label, small_font)
            lx = (x + prev_x) // 2 - lw // 2
            ly = line_y + 16
            draw.text((lx, ly), label, font=small_font, fill=color)

        if step["position"] == "top":
            box = (x - box_w // 2, line_y - circle_r - 24 - box_h, x + box_w // 2, line_y - circle_r - 24)
        else:
            box = (x - box_w // 2, line_y + circle_r + 24, x + box_w // 2, line_y + circle_r + 24 + box_h)

        if is_final:
            draw.rounded_rectangle(box, radius=16, fill=FINAL_BG, outline=box_border, width=2)
            title_color = TEXT_WHITE
            sub_color = (200, 215, 240)
        else:
            draw.rounded_rectangle(box, radius=16, fill=BOX_BG, outline=box_border, width=2)
            title_color = TEXT_DARK
            sub_color = TEXT_MUTED

        draw_centered_text(
            draw,
            box,
            [
                (step["title"], title_font, title_color),
                (step["subtitle"], sub_font, sub_color),
            ],
        )

        draw_circle_node(draw, x, line_y, circle_r, step["num"], step["circle_style"], num_font)

        if step.get("start_label"):
            sw, sh = text_size(draw, step["start_label"], label_font)
            draw.text((x - sw // 2, line_y + circle_r + 8), step["start_label"], font=label_font, fill=C_BLUE)

    header = "فرآیند محاسبه وزن معیارها با روش AHP (چند خبره)"
    hw, _ = text_size(draw, header, title_font)
    draw.text(((width - hw) // 2, 36), header, font=title_font, fill=C_NAVY)

    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, "PNG")


def add_step_slide(prs, step: dict, blank_layout) -> None:
    slide = prs.slides.add_slide(blank_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    num = step["num"]
    title = f"مرحله {num}: {step['title']}"
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(3))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(*C_NAVY)
    p1.alignment = PP_ALIGN.RIGHT

    p2 = tf.add_paragraph()
    p2.text = step["subtitle"]
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(*TEXT_MUTED)
    p2.alignment = PP_ALIGN.RIGHT
    p2.space_before = Pt(16)


def build_pptx(png_path: Path, pptx_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.6))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "فرآیند محاسبه وزن معیارها با روش AHP (چند خبره)"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*C_NAVY)
    p.alignment = PP_ALIGN.CENTER

    slide.shapes.add_picture(str(png_path), Inches(0.2), Inches(0.9), width=Inches(12.9))

    notes_slide = prs.slides.add_slide(blank_layout)
    notes_slide.background.fill.solid()
    notes_slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    notes_title = notes_slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.8))
    ntf = notes_title.text_frame
    np = ntf.paragraphs[0]
    np.text = "شرح مراحل فلوچارت AHP"
    np.font.size = Pt(28)
    np.font.bold = True
    np.font.color.rgb = RGBColor(*C_NAVY)
    np.alignment = PP_ALIGN.RIGHT

    body = notes_slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.5), Inches(5.5))
    btf = body.text_frame
    btf.word_wrap = True
    btf.vertical_anchor = MSO_ANCHOR.TOP

    items = [
        "مرحله ۱ — گردآوری نظرات خبرگان: جمع‌آوری ماتریس‌های مقایسه زوجی از خبرگان",
        "مرحله ۲ — مقایسه‌های زوجی: انجام مقایسات زوجی با روش AHP",
        "مرحله ۳ — بررسی سازگاری: محاسبه نسبت سازگاری (CR) برای هر خبره",
        "مرحله ۴ — حذف شده (پذیرش ماتریس‌ها بر اساس CR < 0.1)",
        "مرحله ۵ — میانگین هندسی: تجمیع نظرات خبرگان",
        "مرحله ۶ — بررسی سازگاری ماتریس تجمیع‌شده: محاسبه CR ماتریس نهایی",
        "مرحله ۷ — وزن نهایی معیارها: استخراج وزن‌ها از ماتریس تجمیع‌شده",
    ]

    for idx, item in enumerate(items):
        para = btf.paragraphs[0] if idx == 0 else btf.add_paragraph()
        para.text = item
        para.font.size = Pt(18)
        para.font.color.rgb = RGBColor(*TEXT_DARK)
        para.alignment = PP_ALIGN.RIGHT
        para.space_after = Pt(10)
        if "حذف شده" in item:
            para.font.color.rgb = RGBColor(180, 60, 60)
            para.font.italic = True

    for step in STEPS:
        add_step_slide(prs, step, blank_layout)

    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(pptx_path))


def main() -> None:
    render_flowchart_png(PNG_PATH)
    build_pptx(PNG_PATH, PPTX_PATH)
    print(f"Created: {PPTX_PATH}")
    print(f"Created: {PNG_PATH}")


if __name__ == "__main__":
    main()
